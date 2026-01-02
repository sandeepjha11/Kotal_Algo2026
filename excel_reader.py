import openpyxl
import config
from logger import logger
import datetime

def load_credentials_and_settings():
    try:
        workbook = openpyxl.load_workbook("trading_system.xlsx")

        # Load Credentials
        credentials_sheet = workbook["Credentials"]
        credentials = list(credentials_sheet.iter_rows(min_row=2, max_row=2, values_only=True))[0]
        config.consumer_key = credentials[0]
        config.CS = credentials[1]
        config.Mob = credentials[2]
        config.Pwd = credentials[3]
        config.MPIN = credentials[4]
        config.totp = credentials[5]
        config.ucc = credentials[6]

        # Load Settings
        settings_sheet = workbook["Settings"]
        settings = list(settings_sheet.iter_rows(min_row=2, max_row=2, values_only=True))[0]
        config.SYMBOL = settings[0]
        try:
            config.QTY = int(settings[1])
        except (ValueError, TypeError):
            logger.warning("Invalid 'Lot Size' value in Settings sheet. Please fill it out.")

        config.EXPIRY_DATE = settings[3]
        entry_time_str = settings[4]
        exit_time_str = settings[5]

        try:
            config.SL_LIMIT = float(settings[6])
        except (ValueError, TypeError):
            logger.warning("Invalid 'SL Limit' value in Settings sheet. Please fill it out.")

        if isinstance(entry_time_str, datetime.time):
            config.ENTRY_TIME = (entry_time_str.hour, entry_time_str.minute, entry_time_str.second)
        elif entry_time_str and 'HH' not in str(entry_time_str):
            try:
                config.ENTRY_TIME = tuple(map(int, str(entry_time_str).split(':')))
            except ValueError:
                logger.error(f"Invalid Entry Time format: {entry_time_str}. Please use HH:MM:SS.")

        if isinstance(exit_time_str, datetime.time):
            config.EXIT_TIME = (exit_time_str.hour, exit_time_str.minute, exit_time_str.second)
        elif exit_time_str and 'HH' not in str(exit_time_str):
            try:
                config.EXIT_TIME = tuple(map(int, str(exit_time_str).split(':')))
            except ValueError:
                logger.error(f"Invalid Exit Time format: {exit_time_str}. Please use HH:MM:SS.")

    except FileNotFoundError:
        logger.error("Error: trading_system.xlsx not found.")
        # Handle error appropriately
    except Exception as e:
        logger.exception(f"An error occurred: {e}")
        # Handle error appropriately

def read_trading_signals(all_trades=False):
    try:
        workbook = openpyxl.load_workbook("trading_system.xlsx")
        trading_sheet = workbook["Trading"]

        signals = []
        for row_index, row in enumerate(trading_sheet.iter_rows(min_row=2, values_only=True), start=2):
            if row[0] is not None:
                if all_trades or row[3] is None: # Check for status is None
                    signals.append({
                        "row_index": row_index,
                        "strike_price": row[0],
                        "option_type": row[1],
                        "buy_sell": row[2],
                        "status": row[3],
                        "entry_price": row[4],
                        "exit_price": row[5],
                        "sl": row[6],
                        "mtm": row[7],
                    })
        return signals
    except FileNotFoundError:
        logger.error("Error: trading_system.xlsx not found.")
        # Handle error appropriately
    except Exception as e:
        logger.exception(f"An error occurred: {e}")
        # Handle error appropriately
    return []
