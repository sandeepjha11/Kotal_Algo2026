import openpyxl

def create_excel_file():
    workbook = openpyxl.Workbook()

    # Credentials Sheet
    credentials_sheet = workbook.active
    credentials_sheet.title = "Credentials"
    credentials_sheet.append(["Consumer Key", "Consumer Secret", "Mobile Number", "Password", "MPIN", "TOTP", "UCC"])
    credentials_sheet.append(["", "", "", "", "", "", ""])
    credentials_sheet.append(["SECURITY WARNING:", "Storing your credentials in a plaintext file is a security risk.", "Anyone with access to this file can access your trading account."])

    # Settings Sheet
    settings_sheet = workbook.create_sheet("Settings")
    settings_sheet.append(["Underlying", "Lot Size", "Premium Amount", "Expiry Dates", "Entry Time", "Exit Time", "SL Limit"])
    settings_sheet.append(["NIFTY", 50, 100, "YYYY-MM-DD", "HH:MM:SS", "HH:MM:SS", 1])

    # Option Chain Sheet
    option_chain_sheet = workbook.create_sheet("Option Chain")
    option_chain_sheet.append(["Strike Price", "LTP", "Open", "High", "Low", "Close", "Volume"])

    # Trading Sheet
    trading_sheet = workbook.create_sheet("Trading")
    trading_sheet.append(["Strike Price", "Option Type", "Buy/Sell", "Status", "Entry Price", "Exit Price", "SL", "MTM"])

    workbook.save("trading_system.xlsx")

def update_option_chain(workbook, option_chain_data):
    option_chain_sheet = workbook["Option Chain"]

    # Clear existing data
    for row in option_chain_sheet.iter_rows(min_row=2):
        for cell in row:
            cell.value = None

    # Add new data
    for option in option_chain_data:
        option_chain_sheet.append(option)

def update_trading_sheet(workbook, row_index, trade_data):
    trading_sheet = workbook["Trading"]

    for col_index, value in enumerate(trade_data, start=1):
        trading_sheet.cell(row=row_index, column=col_index, value=value)

if __name__ == "__main__":
    create_excel_file()
